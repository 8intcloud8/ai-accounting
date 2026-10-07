#!/usr/bin/env python3
"""
inspect_statement.py - Print headers + sample rows of a bank statement file
as JSON, so an agent can reason about what each column means.

Note: this script is deliberately NOT named inspect.py. A file named
inspect.py inside scripts/ (which sits on sys.path when the script runs)
shadows Python's stdlib `inspect` module, breaking anything that
transitively imports it (e.g. pdfplumber) — that was a real bug.

Usage:
    python3 inspect_statement.py --file <path> [--rows 15]
"""
import argparse
import json
import os
import re
import sys

_NUMERIC_RE = re.compile(r"^[\-\+\$\(]?[\d,]+\.?\d*\)?%?$")
_DATE_RE = re.compile(
    r"^\d{1,4}[/\-.]\d{1,2}[/\-.]\d{1,4}$|^\d{1,2}\s+[A-Za-z]{3,9}\s+\d{2,4}$"
)


def _looks_like_data_row(cells):
    """Heuristic: does this row look like transaction data (numbers/dates)
    rather than column names? Used to flag likely headerless CSVs."""
    non_empty = [str(c).strip() for c in cells if c is not None and str(c).strip()]
    if not non_empty:
        return False
    data_like = sum(
        1 for c in non_empty if _NUMERIC_RE.match(c) or _DATE_RE.match(c)
    )
    return data_like >= max(1, len(non_empty) // 2)


def inspect_csv(path, n_rows):
    import csv

    with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
        reader = csv.reader(f)
        rows = list(reader)

    if not rows:
        return {"file": path, "type": "csv", "error": "File is empty."}

    header = rows[0]
    sample = rows[1 : 1 + n_rows]

    result = {
        "file": path,
        "type": "csv",
        "total_rows": len(rows) - 1,
        "header": header,
        "sample_rows": sample,
    }

    if _looks_like_data_row(header):
        result["warning"] = (
            "The first row looks like transaction data (numbers/dates), not column "
            "names — this file may be headerless. If so, set \"header_row\": null in "
            "the mapping and map columns by index instead of name."
        )

    return result


def inspect_excel(path, n_rows):
    try:
        import openpyxl
    except ImportError:
        return {"file": path, "type": "excel", "error": "openpyxl not installed."}

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]

    all_rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        all_rows.append(list(row))
        if i > n_rows + 5:
            break

    if not all_rows:
        return {"file": path, "type": "excel", "error": "Sheet is empty."}

    header = all_rows[0]
    sample = all_rows[1 : 1 + n_rows]
    return {
        "file": path,
        "type": "excel",
        "sheet_name": sheet_name,
        "all_sheets": wb.sheetnames,
        "header": header,
        "sample_rows": sample,
    }


def inspect_pdf(path, n_rows):
    try:
        import pdfplumber
    except ImportError:
        return {"file": path, "type": "pdf", "error": "pdfplumber not installed."}

    result = {"file": path, "type": "pdf", "pages": [], "warning": None}

    total_text_len = 0
    with pdfplumber.open(path) as pdf:
        from nab_pdf import read_listing

        listing = read_listing(pdf)
        if listing is not None:
            header, rows = listing
            return {"file": path, "type": "pdf", "layout": "nab_transaction_listing",
                    "header": header, "total_rows": len(rows),
                    "sample_rows": rows[:n_rows], "balance_status": "reconciled"}
        num_pages = len(pdf.pages)
        for page_idx, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            total_text_len += len(text.strip())

            page_info = {"page": page_idx + 1, "text_length": len(text.strip())}

            tables = page.extract_tables()
            if tables:
                page_info["tables_found"] = len(tables)
                first_table = tables[0]
                page_info["header_guess"] = first_table[0] if first_table else None
                page_info["sample_rows"] = first_table[1 : 1 + n_rows] if len(first_table) > 1 else []
            else:
                page_info["tables_found"] = 0
                lines = [l for l in text.splitlines() if l.strip()]
                page_info["text_lines_sample"] = lines[:n_rows]

            result["pages"].append(page_info)

            if page_idx >= 4:  # cap inspection to first 5 pages for speed
                result["note"] = f"Only first 5 of {num_pages} pages inspected."
                break

    avg_text_per_page = total_text_len / max(len(result["pages"]), 1)
    if avg_text_per_page < 20:
        result["warning"] = (
            "This PDF appears to be SCANNED (little to no extractable text). "
            "OCR would be required before this file can be mapped or extracted. "
            "Do not attempt to guess a column mapping from this output."
        )

    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", required=True)
    parser.add_argument("--rows", type=int, default=15)
    args = parser.parse_args()

    path = args.file
    if not os.path.exists(path):
        print(json.dumps({"file": path, "error": "File not found."}))
        sys.exit(1)

    ext = os.path.splitext(path)[1].lower()

    try:
        if ext == ".csv" or ext == ".tsv":
            result = inspect_csv(path, args.rows)
        elif ext in (".xlsx", ".xlsm", ".xls"):
            result = inspect_excel(path, args.rows)
        elif ext == ".pdf":
            result = inspect_pdf(path, args.rows)
        else:
            result = {"file": path, "error": f"Unsupported file extension: {ext}"}
    except Exception as e:
        result = {"file": path, "error": f"Failed to inspect file: {e}"}

    print(json.dumps(result, indent=2, default=str))


if __name__ == "__main__":
    main()
